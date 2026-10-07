"""Checks that run on every pull request, against the built-in demo database (invented people, no secrets).

They cover the rules the church relies on: follow-up colours, the two-year archive, safe ticks and unticks,
one-winner sign-up approvals, version-checked edits, pastors' lists and the WhatsApp summary.
"""
import datetime as dt

import pytest

import attendance as A


FREE_DAY = (A.today() + dt.timedelta(days=1)).isoformat()  # a date the demo data has no service on


@pytest.fixture()
def store():
    return A.SqlStore(None)  # fresh demo database each time


def test_follow_up_levels_match_the_rules(store):
    df = A.missed_streaks(store.list_members(), store.list_services())
    assert not df.empty
    assert (df[df.missed >= A.RED_AT].level == "red").all()
    assert (df[(df.missed >= A.YELLOW_AT) & (df.missed < A.RED_AT)].level == "yellow").all()
    assert (df[df.missed < A.YELLOW_AT].level == "ok").all()


def test_people_unseen_for_two_years_are_archived_and_come_back_when_ticked(store):
    members, services = store.list_members(), store.list_services()
    old = A.missed_streaks(members, services, archive="only")
    assert set(old.name) == {"Old Friend One", "Old Friend Two", "Old Friend Three"}
    assert not set(old.id) & set(A.missed_streaks(members, services).id)  # left out of follow-up
    back = old.id.iloc[0]
    store.set_present(A.today().isoformat(), back, True)
    assert back not in set(A.missed_streaks(store.list_members(), store.list_services(), archive="only").id)


def test_away_status_takes_someone_off_the_follow_up_list(store):
    df = A.missed_streaks(store.list_members(), store.list_services())
    red = df[df.level == "red"].id.iloc[0]
    store.update_members({red: {"status": "Away"}})
    assert red not in set(A.missed_streaks(store.list_members(), store.list_services()).id)


def test_ticks_are_repeatable_and_stale_unticks_are_blocked(store):
    day, mid = FREE_DAY, store.list_members()[0]["id"]
    assert store.set_present(day, mid, True) == "done"
    assert store.set_present(day, mid, True) == "already"
    assert store.set_present(day, mid, False, seen="2000-01-01T00:00:00+12:00") == "changed"
    assert mid in store.get_service(day)["present"]
    seen = store.get_service(day)["present"][mid]
    assert store.set_present(day, mid, False, seen=seen) == "done"
    assert store.set_present(day, mid, False) == "already"


def test_untick_everyone_clears_one_service_only(store):
    day = FREE_DAY
    ids = [m["id"] for m in store.list_members()[:5]]
    for mid in ids:
        store.set_present(day, mid, True)
    other = sum(len(s["present"]) for s in store.list_services() if s["date"] != day)
    assert store.clear_service(day, by="Test (Admin)") == 5
    assert not store.get_service(day)["present"]
    assert sum(len(s["present"]) for s in store.list_services() if s["date"] != day) == other
    assert any(r["kind"] == "clear_service" for r in store.activity(day))


def test_a_sign_up_can_only_be_approved_once(store):
    reg = store.list_registrations("pending")[0]
    before = len(store.list_members())
    store.approve_registration(reg, None, True)
    with pytest.raises(A.AlreadyHandled):
        store.approve_registration(reg, None, True)
    assert len(store.list_members()) == before + 1


def test_a_failed_approval_leaves_nothing_half_done(store, monkeypatch):
    reg = store.list_registrations("pending")[0]
    before = len(store.list_members())

    def boom(*a, **k):
        raise RuntimeError("simulated failure")
    monkeypatch.setattr(store, "set_present", boom)
    with pytest.raises(RuntimeError):
        store.approve_registration(reg, None, True)
    monkeypatch.undo()
    store._cache.clear()
    assert len(store.list_members()) == before
    assert any(r["id"] == reg["id"] for r in store.list_registrations("pending"))


def test_stale_edits_are_refused(store):
    m = store.list_members()[0]
    seen = {m["id"]: m["version"]}
    assert store.update_members({m["id"]: {"phone": "111"}}, seen) == ([m["id"]], [])
    assert store.update_members({m["id"]: {"phone": "222"}}, seen) == ([], [m["id"]])
    assert {x["id"]: x for x in store.list_members()}[m["id"]]["phone"] == "111"


def test_pastors_have_their_own_lists(store):
    df = A.missed_streaks(store.list_members(), store.list_services())
    sizes = df[df.pastor != ""].groupby("pastor").size()
    assert len(sizes) == 4 and (sizes <= A.PASTOR_GROUP_SIZE).all()
    m = store.list_members()[-1]
    store.update_members({m["id"]: {"pastor": "Pastor Test"}})
    assert {x["id"]: x for x in store.list_members()}[m["id"]]["pastor"] == "Pastor Test"


def test_whatsapp_summary_hides_names_unless_asked(store):
    text = A.whatsapp_summary(store)
    assert text.startswith("*FCC") and "Present:" in text and "Livestream" not in text
    df = A.missed_streaks(store.list_members(), store.list_services())
    red = df[df.level == "red"].name.iloc[0]
    assert red not in text
    assert red in A.whatsapp_summary(store, names=True, link="https://example.com/live")
    assert "https://example.com/live" in A.whatsapp_summary(store, link="https://example.com/live")


def test_whatsapp_messages_can_be_plain_text(store):
    import re
    emoji = re.compile("[\U0001F300-\U0001FAFF\u2600-\u27BF\u2B00-\u2BFF]")
    rows = A.church_numbers(store)
    for plain, rich in ((A.whatsapp_summary(store, link="https://example.com/live", emoji=False),
                         A.whatsapp_summary(store, link="https://example.com/live")),
                        (A.whatsapp_overview(rows, emoji=False), A.whatsapp_overview(rows))):
        assert emoji.search(rich) and not emoji.search(plain)
        assert "Present" in plain and "\n\n" in plain


def test_churches_added_in_the_phone_app_show_up(store):
    assert store.church_names() == []  # no churches table yet: nothing breaks
    store._exec("CREATE TABLE churches (name TEXT PRIMARY KEY)")
    store._exec("INSERT INTO churches (name) VALUES (?)", ("Brisbane",))
    store._churches = None
    assert "Brisbane" in A.all_churches(store)
    assert A.all_churches(store)[0] == A.home_church()


def test_church_list_works_with_a_store_from_before_the_update(store):
    class Old:  # stands in for a connection object created by the previous version of the app
        demo = True
        def __init__(self, real): self._exec, self.list_members = real._exec, real.list_members
    assert A.all_churches(Old(store))[0] == A.home_church()


def test_daily_report_builds(store):
    import report as R
    r = R.build(store, dt.date.today())
    assert r["subject"] and "<html" in r["html"].lower() and len(r["xlsx"]) > 1000


def test_each_church_has_its_own_report_list_and_log(store):
    import report as R
    home, branch = A.ChurchStore(store, A.home_church()), A.ChurchStore(store, "Sydney")
    assert R.recipients(branch) == []  # a branch never borrows another church's list
    assert R.save_recipients(branch, "pastor@sydney.example")[0] == ["pastor@sydney.example"]
    assert R.recipients(branch) == ["pastor@sydney.example"] and "pastor@sydney.example" not in R.recipients(home)
    assert "Sydney" in R.build(branch)["subject"] and A.home_church() in R.build(home)["subject"]
    sent = []
    R._send_brevo = lambda cfg, to, r: sent.append((to, r["subject"]))
    cfg = dict(ready=True, brevo_api_key="test", sender="x@example.com")
    R.send(branch, cfg); R.send(home, cfg); R.send_overview(store, cfg)
    assert [h["kind"] for h in branch.email_history()] == ["manual:Sydney"]
    assert [h["kind"] for h in home.email_history()] == ["manual"]
    empty = A.ChurchStore(store, "Melbourne")
    with pytest.raises(RuntimeError):
        R.send(empty, cfg)


def test_renaming_a_church_moves_everything_with_it(store):
    import report as R
    old = next(c for c in A.all_churches(store) if c != A.home_church())
    people = {m["id"] for m in A.ChurchStore(store, old).list_members()}
    assert people
    R.save_recipients(A.ChurchStore(store, old), "pastor@branch.example")
    assert A.rename_church(store, old, "  New   Name ") == "New Name"
    assert "New Name" in A.all_churches(store) and old not in A.all_churches(store)
    assert {m["id"] for m in A.ChurchStore(store, "New Name").list_members()} == people
    assert R.recipients(A.ChurchStore(store, "New Name")) == ["pastor@branch.example"]
    for bad in (A.home_church(), "Nowhere"):
        with pytest.raises(ValueError):
            A.rename_church(store, bad, "Something")
    with pytest.raises(ValueError):
        A.rename_church(store, "New Name", A.home_church())


def test_the_all_churches_email_has_numbers_but_no_names(store):
    import report as R
    r = R.build_overview(store)
    assert "all churches" in r["subject"] and len(r["xlsx"]) > 1000
    for m in store.list_members():
        assert m["full_name"] not in r["html"] and m["full_name"] not in r["text"]


def test_blue_marks_people_who_missed_the_latest_service_only(store):
    df = A.missed_streaks(store.list_members(), store.list_services())
    assert set(df[df.flag == "blue"].missed) <= {1, 2}
    assert (df[df.flag == "blue"].level == "ok").all()  # blue is only a heads-up, not yet a follow-up call
    assert (df[df.missed == 0].flag == "ok").all()
    assert (df[df.level != "ok"].flag == df[df.level != "ok"].level).all()  # yellow and red are unchanged


def test_adults_and_kids_are_counted_separately(store):
    members = store.list_members()
    mem = {m["id"]: m for m in members}
    kids = [m["id"] for m in members if A.is_child(m)]
    assert kids and len(kids) < len(members)
    assert A.split_ages(list(mem), mem) == (len(members) - len(kids), len(kids))
    assert A.split_ages([kids[0]], mem) == (0, 1)
    store.update_members({kids[0]: {"age_group": "Adult"}})
    assert not A.is_child({x["id"]: x for x in store.list_members()}[kids[0]])
    assert "Adults:" in A.whatsapp_summary(store) and "Average" not in A.whatsapp_summary(store)


def test_follow_up_download_has_the_people_who_need_a_call(store):
    df = A.missed_streaks(store.list_members(), store.list_services())
    need = df[df.level != "ok"]
    table = A.followup_table(need)
    assert list(table.Name) == list(need.name) and {"Status", "Phone", "Pastor", "Adult / Child"} <= set(table.columns)
    assert len(A._xlsx(table, "Needs follow-up")) > 1000


def test_each_church_only_sees_its_own_people(store):
    names = A.all_churches(store)
    assert names[0] == A.home_church() and {"Sydney", "Melbourne"} <= set(names)
    views = {c: A.ChurchStore(store, c) for c in names}
    ids = {c: {m["id"] for m in v.list_members()} for c, v in views.items()}
    assert sum(len(i) for i in ids.values()) == len(store.list_members())  # everyone belongs to exactly one church
    assert not ids["Sydney"] & ids[A.home_church()]
    for c, v in views.items():
        for s in v.list_services():
            assert s["present"] and set(s["present"]) <= ids[c]
    # someone added while looking at Sydney belongs to Sydney, and untick-all there leaves the others alone
    views["Sydney"].upsert_members([dict(id="new-syd", full_name="New Sydney Person", type="member", created_at=A.now_iso())])
    assert "new-syd" in {m["id"] for m in views["Sydney"].list_members()}
    assert "new-syd" not in {m["id"] for m in views[A.home_church()].list_members()}
    home_one, syd_one = next(iter(ids[A.home_church()])), next(iter(ids["Sydney"]))
    store.set_present(FREE_DAY, home_one, True)
    store.set_present(FREE_DAY, syd_one, True)
    assert views["Sydney"].clear_service(FREE_DAY) == 1
    assert home_one in store.get_service(FREE_DAY)["present"]


def test_the_bishops_numbers_contain_no_names(store):
    rows = A.church_numbers(store)
    assert [r["church"] for r in rows] == A.all_churches(store)
    assert sum(r["register"] for r in rows) <= len(store.list_members())
    text = A.whatsapp_overview(rows) + str(rows)
    for m in store.list_members():
        assert m["full_name"] not in text and (not m["phone"] or m["phone"] not in text)


def test_everyone_download_has_a_sheet_per_church(store):
    import io
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(A.all_churches_xlsx(store)))
    assert wb.sheetnames == ["Everyone"] + A.all_churches(store)
    assert wb["Everyone"].max_row - 1 == len(store.list_members())
    assert [c.value for c in wb["Everyone"][1]][0] == "Church"
    moved = A.ChurchStore(store, A.home_church()).list_members()[0]["id"]
    store.update_members({moved: {"church": "Sydney"}})  # HQ moves someone to another church
    assert moved in {m["id"] for m in A.ChurchStore(store, "Sydney").list_members()}


def test_people_per_pastor_can_be_changed_for_each_church(store):
    syd, home = A.ChurchStore(store, "Sydney"), A.ChurchStore(store, A.home_church())
    assert A.group_size(syd) == A.PASTOR_GROUP_SIZE
    syd.set_setting("pastor_group_size:Sydney", "15")
    assert A.group_size(syd) == 15 and A.group_size(home) == A.PASTOR_GROUP_SIZE


def test_passwords_that_are_shared_are_reported(monkeypatch):
    import streamlit as st
    secrets = {"admin_password": "a", "bishop_password": "b", "attendance_password": "t",
               "church_admin_passwords": {"Melbourne": "m-admin"}, "church_passwords": {"Melbourne": "m-team"}}
    monkeypatch.setattr(st, "secrets", secrets)
    assert A.shared_passwords() == []  # the same church name under both headings is fine
    secrets["church_passwords"]["Melbourne"] = "m-admin"
    assert A.shared_passwords() == ["Melbourne church admin and Melbourne team"]


def test_email_addresses_are_checked_in_plain_words():
    import report as R
    for ok in ("a@b.com", "Kelvin.M+x@Gmail.com", "x@church.org.nz"):
        assert R.email_problem(ok) == ""
    assert "space" in R.email_problem("a b@c.com")
    assert "one @" in R.email_problem("a@@b.com")
    assert "after the @" in R.email_problem("a@b")
    assert "did you mean a@gmail.com" in R.email_problem("a@gmial.com")
    assert R.email_problem("")


def test_only_sunday_services_count_as_missed():
    people = [dict(id="a", full_name="Sunday Only", status="", date_joined="2026-01-01"),
              dict(id="b", full_name="Gone Quiet", status="", date_joined="2026-01-01"),
              dict(id="c", full_name="Midweek Only", status="", date_joined="2026-01-01")]
    days = ["2026-09-06", "2026-09-09", "2026-09-13", "2026-09-16", "2026-09-20", "2026-09-23"]  # Sun, Wed, Sun, Wed...
    services = [dict(date=d, present={"a": "t"} if A.is_sunday(d) else {"c": "t"}) for d in days]
    services[0]["present"]["b"] = "t"
    df = A.missed_streaks(people, services, upto=A.dt.date(2026, 9, 24)).set_index("id")
    assert df.loc["a", "missed"] == 0   # never comes on Wednesday, never flagged for it
    assert df.loc["b", "missed"] == 2   # two Sundays missed; the three Wednesdays are not counted
    assert df.loc["c", "missed"] == 0   # seen at the latest midweek service
